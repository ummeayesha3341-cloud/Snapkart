from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Users\Ahmed\OneDrive\Documents\ChatGPT\snapkart\snapkart_self_checkout_survey.docx"
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.75)
sec.left_margin = sec.right_margin = Inches(.8)

for name, size in [("Normal", 10.5), ("Title", 16), ("Heading 1", 12), ("Heading 2", 11)]:
    st = doc.styles[name]
    st.font.name = "Times New Roman"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    st.font.size = Pt(size)
    if name != "Normal":
        st.font.bold = True
        st.font.color.rgb = RGBColor(0,0,0)
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(4)
doc.styles["Normal"].paragraph_format.space_after = Pt(5)
doc.styles["Normal"].paragraph_format.line_spacing = 1.05

def shade(cell, fill):
    e = OxmlElement("w:shd"); e.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(e)

def margins(cell):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement("w:tcMar")
    for n in ("top","start","bottom","end"):
        x = OxmlElement("w:"+n); x.set(qn("w:w"), "100"); x.set(qn("w:type"), "dxa"); m.append(x)
    tcPr.append(m)

def para(text, style=None, center=False):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(text)
    return p

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, v in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=v; shade(c,"1F4E78"); margins(c)
        for r in c.paragraphs[0].runs: r.font.name="Times New Roman"; r.font.size=Pt(8.5); r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    for j, row in enumerate(rows):
        cells=t.add_row().cells
        for i, v in enumerate(row):
            c=cells[i]; c.text=v; margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if j%2: shade(c,"EAF2F8")
            for p in c.paragraphs:
                for r in p.runs: r.font.name="Times New Roman"; r.font.size=Pt(8.3)
    doc.add_paragraph()

para("A Survey of Smart Retail Self Checkout Systems Using Barcode RFID Mobile and Computer Vision Technologies", "Title", True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("SnapKart Survey Paper Draft").italic=True

para("Abstract", "Heading 1")
para("Retail self-checkout technologies aim to reduce checkout delays, improve transaction convenience, and reduce dependency on cashier-operated counters. This survey reviews major smart retail self-checkout approaches: barcode scanning, mobile scan-and-go, RFID-enabled carts, and computer-vision product recognition. The approaches are compared by identification method, infrastructure, cost, computational burden, user participation, and operational risk. Barcode and mobile approaches can build on existing product labels and common camera-enabled devices. RFID can automate identification but requires tags, readers, and supporting infrastructure. Computer vision can help identify visually distinguishable or non-barcoded goods but requires data, model development, and robust in-store imaging. The survey identifies a need for lightweight software-oriented architectures that connect barcode capture, product lookup, digital-cart management, and billing. SnapKart is presented as a conceptual example of this architecture. Future research directions include hybrid recognition, privacy, fraud prevention, accessibility, payment integration, and inventory synchronization.")
para("Keywords: smart retail; self-checkout; barcode recognition; RFID; mobile self-checkout; computer vision; digital cart; retail automation.")

para("I. Introduction", "Heading 1")
para("Retail self-checkout enables shoppers to scan, bag, and sometimes pay for goods with limited direct employee involvement. Its success is not determined by scanning technology alone. Product identification, user interface quality, customer acceptance, privacy, security, staff support, and integration with pricing and inventory systems all affect the checkout experience.")
para("The literature covers fixed self-checkout terminals, mobile applications, RFID smart carts, and camera-driven product-recognition systems. These approaches differ in cost, coverage, reliability, and customer effort. This survey organizes the field into a common taxonomy, compares the approaches using consistent criteria, identifies recurring research gaps, and positions SnapKart as a conceptual lightweight barcode-to-cart framework rather than as a new recognition algorithm.")

para("II. Scope and Review Approach", "Heading 1")
para("This paper is a narrative technology survey. It synthesizes representative peer-reviewed work on retail self-service technology, barcode and mobile checkout, RFID-enabled carts, and computer-vision product recognition. The review focuses on system architecture, identification mechanisms, deployment constraints, service quality, user acceptance, security, and retail operations. The scope is packaged retail goods, with fresh produce treated as a key motivation for vision-based and hybrid systems.")

para("III. Technology Taxonomy", "Heading 1")
para("A. Barcode Based Self Checkout", "Heading 2")
para("Barcode-based checkout uses machine-readable labels such as UPC and EAN codes. A fixed scanner, handheld device, or camera decodes the product identifier and retrieves the corresponding product record. The method is mature and economical because it uses labels already present on many packaged goods. Its weaknesses include damaged labels, poor lighting, unsuitable orientation, and limited coverage for non-barcoded products.")
para("B. Mobile Self Checkout", "Heading 2")
para("Mobile self-checkout transfers scanning from a fixed terminal to a customer smartphone or retailer-provided handheld device. The customer scans items during shopping and maintains a digital cart before payment. Andriulo, Elia, and Gnoni describe mobile self-checkout configurations in fast-moving consumer-goods retail and identify both queue-time potential and scanning constraints [1]. Demoulin and Djelassi show that perceived behavioural control, usefulness, and situational conditions affect use of retail self-service technologies [2].")
para("C. RFID Enabled Smart Carts", "Heading 2")
para("RFID systems identify tagged products using radio-frequency readers. Smart carts may automatically register items and display item details or totals. They reduce the need to present each label to a scanner, but depend on tags, readers, cart hardware, connectivity, and maintenance. Shahroz et al. illustrate the automation potential and infrastructure dependency of RFID smart-cart systems [4].")
para("D. Computer Vision Based Checkout", "Heading 2")
para("Computer-vision checkout uses images or video to identify products. It may combine object detection, classification, weight sensing, and deep neural networks. This is valuable for fresh produce and other goods without consistent barcodes. Hameed, Chai, and Rassau present a coarse-to-fine CNN approach for supermarket fruit and vegetable classification [5], while Nesteruk et al. address adding new checkout classes without manual annotation [6]. These systems broaden product coverage but require representative data, computing resources, and validation under real store conditions.")

para("IV. Comparative Analysis", "Heading 1")
table(["Approach","Identification","Main infrastructure","Primary strength","Key limitation"],[
["Fixed barcode kiosk","Barcode scanner","Terminal and product database","Mature packaged-goods workflow","Centralized queue and manual scan"],
["Mobile self-checkout","Phone camera or handheld scanner","Application, network, database","Flexible and lower hardware cost","User scanning and adoption"],
["RFID smart cart","RFID tags and readers","Tags, readers, cart hardware","Reduced manual scanning","Infrastructure cost"],
["Computer vision","Product appearance and sometimes weight","Camera, dataset, trained model","Can support non-barcoded goods","Data and model burden"],
["Hybrid system","Barcode plus RFID or vision","Integrated hardware and software","Higher coverage and redundancy","Integration complexity"],
])
para("Technology selection should follow the product mix, store format, capital budget, and expected customer journey. Barcode and mobile systems are suitable where packaged goods dominate and low capital expenditure is important. RFID is appropriate where automated identification justifies infrastructure cost. Vision and hybrid approaches are relevant where fresh produce or non-barcoded goods are important.")

para("V. Adoption Service Quality and Customer Experience", "Heading 1")
para("Customer acceptance is a central condition for self-checkout success. Mukerjee, Deshmukh, and Prasad report positive relationships between technology readiness, perceived ease of use, perceived usefulness, and intention to use smartphone self-checkout in an Indian grocery context [3]. Orel and Kara link supermarket self-checkout service quality with customer satisfaction and loyalty [7]. Bulmer, Elms, and Moore also show that self-checkout is embedded in wider social shopping practices, not simply a speed-focused activity [8].")
para("These findings imply that self-checkout should complement, rather than automatically replace, human support. Clear interfaces, exception handling, accessible interaction, and available staff assistance can determine whether a technically functional system produces a satisfactory shopping experience.")

para("VI. Security Privacy and Operations", "Heading 1")
para("Self-checkout creates operational risks that combine physical retail and digital security. Missed scans, product substitution, barcode swapping, payment fraud, and account abuse can create shrinkage exposure. Systems should therefore include transaction logs, secure authentication, exception handling, staff escalation, and proportionate monitoring. Retailers must also protect personal information, purchase histories, device identifiers, and camera-derived data through data minimization, security controls, clear notices, and appropriate retention policies.")
para("Accessibility is equally important. A retail technology that assumes smartphone ownership, perfect camera use, or rapid touchscreen interaction can exclude some shoppers. Assisted modes and staffed alternatives should remain part of inclusive self-checkout design.")

para("VII. Research Gaps", "Heading 1")
table(["Research gap","Why it matters","Promising direction"],[
["Hybrid recognition","Barcodes do not cover all retail goods","Barcode plus vision or weight-assisted verification"],
["Fraud and exceptions","Convenience can increase operational exposure","Privacy-preserving anomaly detection and staff workflows"],
["Accessibility","Users have diverse capabilities and digital skills","Inclusive interface and assisted checkout design"],
["Cost evaluation","Retailers need realistic deployment decisions","Lifecycle-cost studies across store formats"],
["Field validation","Laboratory results may not transfer to stores","Tests across lighting, traffic, and product mixes"],
["Interoperability","Checkout depends on multiple enterprise systems","Standard APIs for catalog, pricing, inventory, and payment"],
])

para("VIII. SnapKart as a Conceptual Software Oriented Framework", "Heading 1")
para("SnapKart illustrates a lightweight software-oriented direction identified by this survey. A camera-enabled device captures a barcode; image processing supports decoding; a backend service retrieves a product record; and a digital cart maintains quantities and the running bill. OpenCV and a barcode-decoding library such as pyzbar can support capture and decoding, while an API layer can connect the client interface with product and cart data.")
para("The framework should be viewed as an integration architecture, not as a new barcode-recognition method. Its potential value is the reuse of existing barcodes and camera-enabled devices without requiring RFID tags, custom carts, or product-specific deep-learning models for packaged goods. Its limitations include unreadable labels, camera and lighting variability, incomplete product catalogs, and limited coverage for fresh or non-barcoded items.")

para("IX. Future Research Directions", "Heading 1")
para("Future work should investigate hybrid barcode and vision systems, fraud prevention, secure payment and account workflows, privacy-preserving monitoring, real-time inventory synchronization, and inclusive interface design. Research is also needed on field deployment cost, staff roles, customer trust, and the effect of self-checkout across diverse retail settings. For image-based systems, robust performance under glare, blur, rotation, occlusion, and crowded baskets remains an important challenge.")

para("X. Conclusion", "Heading 1")
para("Smart retail self-checkout spans fixed barcode terminals, mobile applications, RFID carts, and computer-vision systems. Barcode and mobile approaches provide a practical route for packaged goods because they can rely on existing labels and common devices. RFID offers greater automation at the cost of added infrastructure. Computer vision expands coverage to non-barcoded goods but requires data, models, and dependable in-store operation. The survey shows that successful self-checkout depends on technology, customer acceptance, service quality, security, privacy, accessibility, and integration with retail operations. SnapKart is a conceptual barcode-to-cart framework that can be investigated further through real-world evaluation.")

para("References", "Heading 1")
refs=[
"[1] S. Andriulo, V. Elia, and M. G. Gnoni, \"Mobile self-checkout systems in the FMCG retail sector: A comparison analysis,\" International Journal of RF Technologies, vol. 6, no. 4, pp. 207-224, 2015, doi: 10.3233/RFT-150067.",
"[2] N. T. M. Demoulin and S. Djelassi, \"An integrated model of self-service technology usage in a retail context,\" International Journal of Retail and Distribution Management, vol. 44, no. 5, pp. 540-559, 2016, doi: 10.1108/IJRDM-08-2015-0122.",
"[3] H. S. Mukerjee, G. K. Deshmukh, and U. D. Prasad, \"Technology readiness and likelihood to use self-checkout services using smartphone in retail grocery stores: Empirical evidences from Hyderabad, India,\" Business Perspectives and Research, vol. 7, no. 1, pp. 1-15, 2019, doi: 10.1177/2278533718800118.",
"[4] M. Shahroz, M. F. Mushtaq, M. Ahmad, S. Ullah, A. Mehmood, and G. S. Choi, \"IoT based smart shopping cart using radio frequency identification,\" IEEE Access, vol. 8, pp. 68426-68438, 2020, doi: 10.1109/ACCESS.2020.2986681.",
"[5] K. Hameed, D. Chai, and A. Rassau, \"A sample weight and AdaBoost CNN-based coarse to fine classification of fruit and vegetables at a supermarket self-checkout,\" Applied Sciences, vol. 10, no. 23, Art. no. 8667, 2020, doi: 10.3390/app10238667.",
"[6] S. Nesteruk et al., \"PseudoAugment: Enabling smart checkout adoption for new classes without human annotation,\" IEEE Access, vol. 11, pp. 76869-76882, 2023, doi: 10.1109/ACCESS.2023.3296854.",
"[7] F. D. Orel and A. Kara, \"Supermarket self-checkout service quality, customer satisfaction, and loyalty: Empirical evidence from an emerging market,\" Journal of Retailing and Consumer Services, vol. 21, no. 2, pp. 118-129, 2014, doi: 10.1016/j.jretconser.2013.07.002.",
"[8] S. Bulmer, J. Elms, and S. Moore, \"Exploring the adoption of self-service checkouts and the associated social obligations of shopping practices,\" Journal of Retailing and Consumer Services, vol. 42, pp. 107-116, 2018, doi: 10.1016/j.jretconser.2018.01.016.",
"[9] P. Sharma, V. Kumar, and R. P. J. Kingshott, \"Self-service technology in supermarkets - Do frontline staff still matter?\" Journal of Retailing and Consumer Services, vol. 59, Art. no. 102356, 2021, doi: 10.1016/j.jretconser.2020.102356."
]
for r in refs: para(r)

doc.save(OUT)
